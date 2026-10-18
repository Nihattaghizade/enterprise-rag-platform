import io
from typing import Optional

from bs4 import BeautifulSoup
from docx import Document as DocxDocument
from pptx import Presentation
from pypdf import PdfReader

from app.core.errors import ValidationError
from app.ingestion.types import ParsedDocument

SUPPORTED_MIME_TYPES = {
    "text/plain",
    "text/markdown",
    "text/html",
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}


def parse_docx(raw_bytes: bytes) -> ParsedDocument:
    """Extract text and heading structure from a DOCX file."""
    doc = DocxDocument(io.BytesIO(raw_bytes))
    lines: list[str] = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style_name = (para.style.name if para.style else "").lower()
        if style_name.startswith("heading 1"):
            lines.append(f"# {text}")
        elif style_name.startswith("heading 2"):
            lines.append(f"## {text}")
        elif style_name.startswith("heading"):
            lines.append(f"### {text}")
        else:
            lines.append(text)

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))

    return ParsedDocument(
        text="\n\n".join(lines),
        mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )


def parse_pptx(raw_bytes: bytes) -> ParsedDocument:
    """Extract slide titles and text frames from a PPTX presentation."""
    prs = Presentation(io.BytesIO(raw_bytes))
    slide_texts: list[str] = []
    for idx, slide in enumerate(prs.slides, start=1):
        parts: list[str] = []
        title_text: Optional[str] = None
        if slide.shapes.title and slide.shapes.title.text.strip():
            title_text = slide.shapes.title.text.strip()
            parts.append(f"# {idx}. {title_text}")
        else:
            parts.append(f"# {idx}. Slide {idx}")

        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if slide.shapes.title and shape == slide.shapes.title:
                continue
            frame_text = shape.text_frame.text.strip()
            if frame_text:
                parts.append(frame_text)

        slide_texts.append("\n".join(parts))

    return ParsedDocument(
        text="\n\n".join(slide_texts),
        mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        page_count=len(prs.slides),
    )


def parse_pdf(raw_bytes: bytes) -> ParsedDocument:
    """Extract text from a PDF document (text-only, no OCR per Section 2)."""
    reader = PdfReader(io.BytesIO(raw_bytes))
    page_texts: list[str] = []
    for page in reader.pages:
        extracted = page.extract_text() or ""
        if extracted.strip():
            page_texts.append(extracted.strip())

    return ParsedDocument(
        text="\n\n".join(page_texts),
        mime_type="application/pdf",
        page_count=len(reader.pages),
    )


def parse_html(raw_bytes: bytes) -> ParsedDocument:
    """Extract readable text and headings from an HTML document."""
    decoded = raw_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(decoded, "html.parser")

    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()

    lines: list[str] = []
    for element in soup.find_all(["h1", "h2", "h3", "h4", "p", "li"]):
        content = element.get_text(" ", strip=True)
        if not content:
            continue
        if element.name == "h1":
            lines.append(f"# {content}")
        elif element.name == "h2":
            lines.append(f"## {content}")
        elif element.name in ("h3", "h4"):
            lines.append(f"### {content}")
        else:
            lines.append(content)

    if not lines:
        fallback = soup.get_text("\n", strip=True)
        if fallback:
            lines.append(fallback)

    return ParsedDocument(
        text="\n\n".join(lines),
        mime_type="text/html",
    )


def parse_file(
    raw_bytes: bytes,
    mime_type: str,
    filename: Optional[str] = None,
) -> ParsedDocument:
    """Dispatch file extraction based on MIME type or file extension."""
    normalized_mime = (mime_type or "").split(";")[0].strip().lower()
    lower_name = (filename or "").lower()

    if normalized_mime == "application/pdf" or lower_name.endswith(".pdf"):
        return parse_pdf(raw_bytes)
    if (
        normalized_mime
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or lower_name.endswith(".docx")
    ):
        return parse_docx(raw_bytes)
    if (
        normalized_mime
        == "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        or lower_name.endswith(".pptx")
    ):
        return parse_pptx(raw_bytes)
    if normalized_mime == "text/html" or lower_name.endswith((".html", ".htm")):
        return parse_html(raw_bytes)
    if (
        normalized_mime in ("text/plain", "text/markdown")
        or lower_name.endswith((".txt", ".md"))
    ):
        text = raw_bytes.decode("utf-8", errors="replace")
        resolved_mime = "text/markdown" if lower_name.endswith(".md") else "text/plain"
        return ParsedDocument(text=text, mime_type=resolved_mime)

    raise ValidationError(
        f"Unsupported document format: '{mime_type}'.",
        code="unsupported_format",
    )