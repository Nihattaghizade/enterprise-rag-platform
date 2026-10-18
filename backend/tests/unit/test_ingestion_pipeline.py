import io

import pytest
from docx import Document as DocxDocument
from pptx import Presentation

from app.core.errors import ValidationError
from app.ingestion.chunking import build_chunk_plan
from app.ingestion.normalize import normalize_text, validate_minimum_text
from app.ingestion.parser import parse_file
from app.ingestion.sections import split_into_sections


def test_parse_txt_md_html_docx_pptx() -> None:
    # TXT & MD
    txt_parsed = parse_file(b"# 1. Intro\n\nEnterprise RAG platform overview text.", "text/plain", "doc.txt")
    assert "Enterprise RAG platform" in txt_parsed.text

    # HTML
    html_bytes = b"<html><body><h1>1. Security Policy</h1><p>All queries must filter by tenant_id.</p></body></html>"
    html_parsed = parse_file(html_bytes, "text/html", "policy.html")
    assert "# 1. Security Policy" in html_parsed.text
    assert "tenant_id" in html_parsed.text

    # DOCX
    docx_buf = io.BytesIO()
    doc = DocxDocument()
    doc.add_heading("1. Access Control", level=1)
    doc.add_paragraph("Every tenant is isolated at the database level.")
    doc.save(docx_buf)
    docx_parsed = parse_file(
        docx_buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "access.docx",
    )
    assert "# 1. Access Control" in docx_parsed.text

    # PPTX
    pptx_buf = io.BytesIO()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Architecture Overview"
    prs.save(pptx_buf)
    pptx_parsed = parse_file(
        pptx_buf.getvalue(),
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "deck.pptx",
    )
    assert "Architecture Overview" in pptx_parsed.text
    assert pptx_parsed.page_count == 1


def test_normalize_and_minimum_text_validation() -> None:
    raw = "#   1.1   Heading Title  \r\n\r\n\r\nParagraph   with   extra   spaces."
    normalized = normalize_text(raw)
    assert normalized == "# 1.1 Heading Title\n\nParagraph with extra spaces."

    with pytest.raises(ValidationError):
        validate_minimum_text("short")


def test_sections_and_recursive_vs_semantic_chunking_policy() -> None:
    # Short section (<= 2,500 tokens) -> recursive only
    short_doc = "# 1. Overview\n\n" + ("This is a standard policy sentence for recursive chunking. " * 40)
    normalized_short = normalize_text(short_doc)
    short_sections = split_into_sections(normalized_short)
    assert len(short_sections) == 1
    assert short_sections[0].section_number == "1"

    short_plan = build_chunk_plan(short_sections)
    strategies = {c.strategy for c in short_plan.children}
    assert strategies == {"recursive"}

    # Long section (> 2,500 tokens) -> BOTH recursive (fallback) and semantic (primary)
    long_doc = "# 2. Detailed Specification\n\n" + (
        "Security isolation requires server-derived tenant_id filtering across PostgreSQL and Qdrant. " * 260
    )
    normalized_long = normalize_text(long_doc)
    long_sections = split_into_sections(normalized_long)
    assert long_sections[0].token_count > 2500

    long_plan = build_chunk_plan(long_sections)
    long_strategies = {c.strategy for c in long_plan.children}
    assert "recursive" in long_strategies
    assert "semantic" in long_strategies
    for child in long_plan.children:
        assert child.token_count <= 400